"""Recover original email workbooks and compare them with the delivered Landsat CSV.

Read-only API calls; original attachments stay in ignored local_archive/.
"""
import base64
import csv
import hashlib
import json
import subprocess
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
CACHE = ROOT / "local_archive/metadata_audit_2026-10-16"
OUT = ROOT / "data/external/source_metadata_2026-10-16"
SOURCES = [("1731158abac79364", "MPA_Sites_Controls.xlsx"),
           ("183fa8c993c82b45", "MPA_Ref_Locations.xlsx")]


def api(*parts, params):
    result = subprocess.run(["gws", *parts, "--params", json.dumps(params)],
                            capture_output=True, check=True, text=True)
    return json.loads(result.stdout)


def flatten(part):
    yield part
    for child in part.get("parts", []):
        yield from flatten(child)


def recover(mid, filename):
    path = CACHE / filename
    if not path.exists():
        message = api("gmail", "users", "messages", "get",
                      params={"userId": "me", "id": mid, "format": "full"})
        part = next(p for p in flatten(message["payload"]) if p.get("filename") == filename)
        data = api("gmail", "users", "messages", "attachments", "get",
                   params={"userId": "me", "messageId": mid, "id": part["body"]["attachmentId"]})
        encoded = data["data"]
        path.write_bytes(base64.urlsafe_b64decode(encoded + "=" * (-len(encoded) % 4)))
    return path


def main():
    CACHE.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    raw = ROOT / "data/LANDSAT/MPA_Runs_new.csv"
    with raw.open(encoding="utf-8-sig", newline="") as handle:
        records = list(csv.DictReader(handle))
    years = [key for key in records[0] if key.isdigit()]
    original = np.array([[float(row[y]) for y in years] for row in records])
    report = {"input_sha256": hashlib.sha256(raw.read_bytes()).hexdigest(),
              "input_rows": len(records), "input_years": [years[0], years[-1]], "workbooks": []}
    for mid, filename in SOURCES:
        path = recover(mid, filename)
        wb = load_workbook(path, read_only=True, data_only=True)
        item = {"message_id": mid, "filename": filename,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "sheets": []}
        for sheet in wb:
            values = list(sheet.values)
            numeric_matches = []
            # Search complete annual vectors, not coincidental single-cell matches.
            for r, row in enumerate(values):
                for start in range(max(0, len(row) - len(years) + 1)):
                    window = row[start:start + len(years)]
                    if sum(isinstance(x, (float, int)) for x in window) < 30:
                        continue
                    try:
                        vector = np.array([float(x) if x is not None else np.nan for x in window])
                    except (ValueError, TypeError):
                        continue
                    matches = np.flatnonzero(np.all(np.isclose(original, vector, rtol=0, atol=0.000001,
                                                               equal_nan=True), axis=1))
                    for idx in matches:
                        numeric_matches.append({"excel_row": r + 1, "start_column": start + 1,
                                                "input_row": int(idx) + 2,
                                                "MPA": records[idx]["MPA"],
                                                "status": records[idx]["status"], "rep": records[idx]["rep"]})
            item["sheets"].append({"name": sheet.title, "rows": sheet.max_row,
                                    "columns": sheet.max_column,
                                    "preview": [[str(x) if x is not None else "" for x in row[:12]]
                                                for row in values[:4]],
                                    "complete_annual_vector_matches": numeric_matches})
            if "Max_Annual" in sheet.title:
                usable = [(int(row[0]), row) for row in values[1:]
                          if isinstance(row[0], (int, float)) and str(int(row[0])) in years]
                partial = []
                for col, site in enumerate(values[0][1:], 1):
                    vector = np.array([float(row[col]) if row[col] is not None else np.nan
                                       for _, row in usable])
                    original_subset = original[:, [years.index(str(y)) for y, _ in usable]]
                    matches = np.flatnonzero(np.all(np.isclose(original_subset, vector, rtol=0,
                                                               atol=0.000001, equal_nan=True), axis=1))
                    for idx in matches:
                        partial.append({"site_number": site, "input_row": int(idx) + 2,
                                        "MPA": records[idx]["MPA"], "status": records[idx]["status"],
                                        "years_compared": len(usable)})
                item["sheets"][-1]["common_year_vector_matches"] = partial
        wb.close()
        report["workbooks"].append(item)
    empty_path = OUT / "satellite_empty_reference_rows.csv"
    if empty_path.exists():
        with empty_path.open() as handle:
            empty = [r for r in csv.DictReader(handle) if r["all_delivered_years_zero"] == "TRUE"]
        comparison = []
        for flagged in empty:
            reference = [r for r in records if r["MPA"] == flagged["MPA"] and r["status"] == "reference"]
            retained = [r for r in reference if r["rep"] != flagged["rep"]]
            for year in years:
                a = np.array([float(r[year]) for r in reference])
                b = np.array([float(r[year]) for r in retained])
                current = float(a[np.isfinite(a)].mean()) if np.isfinite(a).any() else None
                candidate = float(b[np.isfinite(b)].mean()) if np.isfinite(b).any() else None
                comparison.append({"MPA": flagged["MPA"], "year": year,
                                   "delivered_reference_mean": current,
                                   "without_candidate_empty_slot": candidate,
                                   "candidate_over_delivered": candidate / current
                                   if candidate is not None and current else None})
        with (OUT / "satellite_empty_slot_arithmetic.csv").open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(comparison[0]))
            writer.writeheader()
            writer.writerows(comparison)
        report["candidate_empty_slot_check"] = {
            "MPAs": [r["MPA"] for r in empty],
            "years_delivered_zero_but_candidate_missing": sum(
                r["delivered_reference_mean"] == 0 and r["without_candidate_empty_slot"] is None
                for r in comparison),
            "nonzero_reference_multipliers": sorted(set(
                r["candidate_over_delivered"] for r in comparison if r["candidate_over_delivered"] is not None)),
            "scope": "Descriptive denominator check only; primary inputs and models unchanged."}
    (OUT / "workbook_comparison.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report.get("candidate_empty_slot_check", {}), indent=2))
    print(json.dumps({"input_rows": report["input_rows"], "workbooks": [
        {"filename": w["filename"], "sheets": [{"name": s["name"],
         "complete_matches": len(s["complete_annual_vector_matches"]),
         "common_year_matches": len(s.get("common_year_vector_matches", []))}
         for s in w["sheets"]]} for w in report["workbooks"]]}, indent=2))


if __name__ == "__main__":
    main()
